"""
Genera la presentacion de PowerPoint de Virthub 2.0.

Diseno: 16:9, fondo oscuro con acento verde azulado (el mismo del sitio), tarjetas
y diagramas hechos con formas editables.
"""

import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
IMG = os.path.join(RAIZ, "docs", "img")
SALIDA = os.path.join(RAIZ, "docs", "Presentacion-Virthub-2.0.pptx")

# Paleta
FONDO = RGBColor(0x0C, 0x1A, 0x18)
FONDO_CARTA = RGBColor(0x13, 0x2A, 0x26)
VERDE = RGBColor(0x3F, 0xC1, 0x9F)
VERDE_OSCURO = RGBColor(0x1F, 0x6F, 0x5C)
BLANCO = RGBColor(0xF2, 0xF7, 0xF5)
GRIS = RGBColor(0x9F, 0xB4, 0xAE)
AMBAR = RGBColor(0xE8, 0xA3, 0x3D)
AZUL = RGBColor(0x4A, 0x9F, 0xE0)
ROJO = RGBColor(0xE0, 0x6C, 0x5E)

FUENTE = "Segoe UI"
FUENTE_TITULO = "Segoe UI Semibold"

ANCHO = Inches(13.333)
ALTO = Inches(7.5)

prs = Presentation()
prs.slide_width = ANCHO
prs.slide_height = ALTO

VACIA = prs.slide_layouts[6]


# ---------------------------------------------------------------------------
# Utilidades
# ---------------------------------------------------------------------------


def nueva_diapositiva(fondo: RGBColor = FONDO):
    slide = prs.slides.add_slide(VACIA)
    fondo_forma = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, ANCHO, ALTO)
    fondo_forma.fill.solid()
    fondo_forma.fill.fore_color.rgb = fondo
    fondo_forma.line.fill.background()
    fondo_forma.shadow.inherit = False
    return slide


def texto(slide, izq, arriba, ancho, alto, contenido, tam=18, color=BLANCO,
          negrita=False, alineacion=PP_ALIGN.LEFT, fuente=FUENTE, espaciado=6,
          anclaje=MSO_ANCHOR.TOP):
    caja = slide.shapes.add_textbox(Inches(izq), Inches(arriba), Inches(ancho), Inches(alto))
    marco = caja.text_frame
    marco.word_wrap = True
    marco.vertical_anchor = anclaje

    if isinstance(contenido, str):
        contenido = [contenido]

    for i, linea in enumerate(contenido):
        parrafo = marco.paragraphs[0] if i == 0 else marco.add_paragraph()
        parrafo.alignment = alineacion
        parrafo.space_after = Pt(espaciado)
        run = parrafo.add_run()
        run.text = linea
        run.font.size = Pt(tam)
        run.font.bold = negrita
        run.font.color.rgb = color
        run.font.name = fuente
    return caja


def titulo_diapositiva(slide, titulo, subtitulo=None, numero=None):
    barra = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.75), Inches(0.55), Inches(0.09), Inches(0.62))
    barra.fill.solid()
    barra.fill.fore_color.rgb = VERDE
    barra.line.fill.background()

    texto(slide, 1.05, 0.48, 11.2, 0.8, titulo, tam=30, color=BLANCO, negrita=True, fuente=FUENTE_TITULO)

    if subtitulo:
        texto(slide, 1.05, 1.16, 11.2, 0.45, subtitulo, tam=14, color=GRIS)

    if numero is not None:
        texto(slide, 12.1, 6.85, 1.0, 0.4, str(numero), tam=11, color=GRIS, alineacion=PP_ALIGN.RIGHT)


def vinietas(slide, elementos, izq=1.05, arriba=1.85, ancho=11.3, tam=17, alto=4.6, color=BLANCO):
    caja = slide.shapes.add_textbox(Inches(izq), Inches(arriba), Inches(ancho), Inches(alto))
    marco = caja.text_frame
    marco.word_wrap = True

    for i, (texto_linea, nivel) in enumerate(elementos):
        parrafo = marco.paragraphs[0] if i == 0 else marco.add_paragraph()
        parrafo.level = nivel
        parrafo.space_after = Pt(10 if nivel == 0 else 6)

        marca = "▸  " if nivel == 0 else "–  "
        run = parrafo.add_run()
        run.text = marca
        run.font.size = Pt(tam - (2 if nivel else 0))
        run.font.color.rgb = VERDE if nivel == 0 else GRIS
        run.font.name = FUENTE

        run2 = parrafo.add_run()
        run2.text = texto_linea
        run2.font.size = Pt(tam - (2 if nivel else 0))
        run2.font.color.rgb = color if nivel == 0 else GRIS
        run2.font.name = FUENTE
    return caja


def tarjeta(slide, izq, arriba, ancho, alto, titulo, cuerpo, acento=VERDE, tam_titulo=16, tam_cuerpo=12):
    forma = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(izq), Inches(arriba), Inches(ancho), Inches(alto))
    forma.fill.solid()
    forma.fill.fore_color.rgb = FONDO_CARTA
    forma.line.color.rgb = acento
    forma.line.width = Pt(1.25)
    forma.shadow.inherit = False
    forma.adjustments[0] = 0.08

    texto(slide, izq + 0.22, arriba + 0.18, ancho - 0.44, 0.5, titulo, tam=tam_titulo, color=acento, negrita=True)
    texto(slide, izq + 0.22, arriba + 0.72, ancho - 0.44, alto - 0.9, cuerpo, tam=tam_cuerpo, color=BLANCO, espaciado=4)
    return forma


def flecha(slide, izq, arriba, ancho, alto=0.35, color=VERDE):
    forma = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(izq), Inches(arriba), Inches(ancho), Inches(alto))
    forma.fill.solid()
    forma.fill.fore_color.rgb = color
    forma.line.fill.background()
    forma.shadow.inherit = False


def imagen_ajustada(slide, ruta, izq, arriba, ancho_max, alto_max, borde=True):
    if not os.path.exists(ruta):
        return None

    from PIL import Image

    with Image.open(ruta) as im:
        proporcion = im.width / im.height

    ancho = ancho_max
    alto = ancho / proporcion

    if alto > alto_max:
        alto = alto_max
        ancho = alto * proporcion

    izq_final = izq + (ancho_max - ancho) / 2
    arriba_final = arriba + (alto_max - alto) / 2

    imagen = slide.shapes.add_picture(ruta, Inches(izq_final), Inches(arriba_final), Inches(ancho), Inches(alto))

    if borde:
        marco = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(izq_final - 0.04), Inches(arriba_final - 0.04),
            Inches(ancho + 0.08), Inches(alto + 0.08),
        )
        marco.fill.background()
        marco.line.color.rgb = VERDE_OSCURO
        marco.line.width = Pt(1)
        marco.shadow.inherit = False
        marco.adjustments[0] = 0.03
        # Mandar al fondo relativo: se reordena moviendo el marco antes que la imagen.
        sp = marco._element
        sp.getparent().remove(sp)
        imagen._element.addprevious(sp)

    return imagen


def pie(slide, texto_pie):
    texto(slide, 0.75, 6.95, 11.2, 0.4, texto_pie, tam=10, color=GRIS)


def separador(slide, kicker, titulo, cuerpo=None):
    texto(slide, 0.95, 2.35, 11.4, 0.5, kicker.upper(), tam=13, color=VERDE, negrita=True)
    texto(slide, 0.95, 2.75, 11.4, 1.0, titulo, tam=38, color=BLANCO, negrita=True, fuente=FUENTE_TITULO)
    if cuerpo:
        texto(slide, 0.95, 3.9, 10.6, 1.2, cuerpo, tam=16, color=GRIS)


# ---------------------------------------------------------------------------
# 1. Portada
# ---------------------------------------------------------------------------

slide = nueva_diapositiva(RGBColor(0x08, 0x14, 0x12))

# Banda de acento
banda = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.22), ALTO)
banda.fill.solid()
banda.fill.fore_color.rgb = VERDE
banda.line.fill.background()

texto(slide, 1.1, 1.55, 11.0, 0.5, "UNIVERSIDAD TECNOLÓGICA DE TULA-TEPEJI", tam=13, color=VERDE, negrita=True)
texto(slide, 1.1, 2.05, 11.0, 1.5, "Virthub 2.0", tam=76, color=BLANCO, negrita=True, fuente=FUENTE_TITULO)
texto(
    slide, 1.1, 3.55, 10.2, 1.3,
    ["Escritorios virtuales, foro y chat para estudiantes",
     "que no pueden cambiar de computadora"],
    tam=24, color=VERDE, espaciado=2,
)

linea = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.1), Inches(5.05), Inches(3.2), Inches(0.05))
linea.fill.solid()
linea.fill.fore_color.rgb = VERDE_OSCURO
linea.line.fill.background()

texto(
    slide, 1.1, 5.35, 11.0, 1.2,
    ["Proyecto integrador  ·  Presentación técnica",
     "Francisco González  ·  Octubre 2026"],
    tam=15, color=GRIS, espaciado=4,
)

# ---------------------------------------------------------------------------
# 2. El problema
# ---------------------------------------------------------------------------

slide = nueva_diapositiva()
titulo_diapositiva(slide, "El problema", "Lo que nos encontramos en el aula", 2)

tarjeta(slide, 0.95, 1.95, 3.7, 3.5, "Equipos viejos",
        ["Computadoras de hace más de diez años.",
         "2 a 4 GB de memoria.",
         "Procesadores Core 2 Duo o Atom."], acento=ROJO, tam_cuerpo=14)

tarjeta(slide, 4.85, 1.95, 3.7, 3.5, "Programas que no corren",
        ["Las prácticas piden herramientas pesadas.",
         "No hay presupuesto para renovar equipos.",
         "El estudiante se queda atrás por el hardware."], acento=AMBAR, tam_cuerpo=14)

tarjeta(slide, 8.75, 1.95, 3.7, 3.5, "Consecuencia",
        ["Se abandona la práctica.",
         "Se pierde el curso.",
         "Y a veces, la carrera."], acento=AZUL, tam_cuerpo=14)

texto(slide, 0.95, 5.75, 11.5, 0.9,
      "El talento no debería depender del equipo que cada quien puede pagar.",
      tam=22, color=VERDE, negrita=True, alineacion=PP_ALIGN.CENTER)

# ---------------------------------------------------------------------------
# 3. La idea
# ---------------------------------------------------------------------------

slide = nueva_diapositiva()
titulo_diapositiva(slide, "La idea", "¿Y si el trabajo pesado no ocurre en tu computadora?", 3)

texto(slide, 0.95, 1.9, 11.5, 1.0,
      "Virthub lleva la universidad al navegador. El equipo solo dibuja la pantalla.",
      tam=20, color=BLANCO)

vinietas(slide, [
    ("Todo el procesamiento ocurre en el servidor.", 0),
    ("El estudiante usa programas actuales desde cualquier equipo.", 0),
    ("La misma cuenta sirve para el foro, el chat y el escritorio.", 0),
    ("Funciona en Windows antiguo, Linux ligero o un equipo prestado.", 0),
    ("Sin instalar nada: solo un navegador.", 0),
], arriba=2.9, tam=19, alto=3.6)

# ---------------------------------------------------------------------------
# 4. Arquitectura
# ---------------------------------------------------------------------------

slide = nueva_diapositiva()
titulo_diapositiva(slide, "Arquitectura", "Servicios independientes, todos en Docker", 4)

# Navegador
tarjeta(slide, 0.6, 2.15, 2.3, 1.5, "Navegador", ["Solo dibuja", "la pantalla"], acento=GRIS, tam_titulo=14, tam_cuerpo=12)

flecha(slide, 3.05, 2.72, 0.75)

# Caddy
tarjeta(slide, 3.9, 2.15, 2.5, 1.5, "Caddy", ["Proxy inverso", "Control de acceso"], acento=VERDE, tam_titulo=14, tam_cuerpo=12)

flecha(slide, 6.5, 2.72, 0.75)

# App
tarjeta(slide, 7.35, 2.15, 2.6, 1.5, "Aplicación", ["Laravel 13", "PHP 8.4"], acento=VERDE, tam_titulo=14, tam_cuerpo=12)

# Servicios
tarjeta(slide, 0.95, 4.35, 2.6, 1.6, "MariaDB", ["Usuarios", "Publicaciones"], acento=AZUL, tam_titulo=14, tam_cuerpo=11)
tarjeta(slide, 3.85, 4.35, 2.6, 1.6, "Sherlock", ["Moderación", "de imágenes"], acento=AMBAR, tam_titulo=14, tam_cuerpo=11)
tarjeta(slide, 6.75, 4.35, 2.6, 1.6, "ClamAV", ["Antivirus", "de archivos"], acento=ROJO, tam_titulo=14, tam_cuerpo=11)
tarjeta(slide, 9.65, 4.35, 2.6, 1.6, "Escritorios", ["Webtop", "ct0 … ctN"], acento=VERDE, tam_titulo=14, tam_cuerpo=11)

texto(slide, 0.95, 6.35, 11.5, 0.5,
      "Cada servicio se puede reiniciar sin tumbar los demás.",
      tam=13, color=GRIS)

# ---------------------------------------------------------------------------
# 5. Stack
# ---------------------------------------------------------------------------

slide = nueva_diapositiva()
titulo_diapositiva(slide, "Con qué está hecho", "Tecnologías del proyecto", 5)

vinietas(slide, [
    ("PHP 8.4 con Laravel 13: toda la lógica de la aplicación.", 0),
    ("MariaDB 11.4: cuentas, publicaciones y reportes.", 0),
    ("Docker: cada pieza en su contenedor.", 0),
    ("Caddy: proxy inverso y control de acceso.", 0),
    ("Cloudflare Tunnel: publica el sitio sin abrir puertos.", 0),
    ("Nginx y PHP-FPM dentro del contenedor de la aplicación.", 0),
    ("Ollama: modelo de lenguaje local para el asistente.", 0),
    ("Python: los servicios de moderación de imágenes.", 0),
], arriba=1.95, tam=18, alto=4.7)

# ---------------------------------------------------------------------------
# 6. Separador: escritorios
# ---------------------------------------------------------------------------

slide = nueva_diapositiva()
separador(slide, "Parte 1", "Escritorios virtuales",
          "Un contenedor por estudiante, entregado por el proxy solo a su dueño.")

# ---------------------------------------------------------------------------
# 7. Webtop
# ---------------------------------------------------------------------------

slide = nueva_diapositiva()
titulo_diapositiva(slide, "Los escritorios Webtop", "Un contenedor por persona", 7)

vinietas(slide, [
    ("Cada escritorio es un contenedor Linux independiente: ct0, ct1, ct2…", 0),
    ("Lleva escritorio Ubuntu MATE, ofimática, navegador y editor.", 0),
    ("La aplicación decide qué contenedor le toca a cada cuenta.", 0),
    ("El navegador solo recibe la imagen: el trabajo ocurre en el servidor.", 0),
    ("El administrador puede probar cada escritorio desde el panel.", 0),
], arriba=1.9, tam=17, alto=3.2)

imagen_ajustada(slide, os.path.join(IMG, "14-escritorio-webtop.png"), 1.2, 4.35, 10.9, 2.35)

# ---------------------------------------------------------------------------
# 8. Tokenizacion: el problema
# ---------------------------------------------------------------------------

slide = nueva_diapositiva()
titulo_diapositiva(slide, "El problema de acceso", "Por qué no bastaba la cookie", 8)

tarjeta(slide, 0.95, 1.95, 5.4, 2.2, "Lo que pasaba",
        ["La comprobación dependía de la cookie de sesión.",
         "La cookie no siempre viaja al subdominio del escritorio.",
         "Resultado: 403 al abrir tu propio escritorio."],
        acento=ROJO, tam_cuerpo=14)

tarjeta(slide, 6.9, 1.95, 5.5, 2.2, "Lo que no se podía hacer",
        ["Quitar la comprobación.",
         "Porque entonces cualquiera con la dirección entra.",
         "Incluido el escritorio de otro."],
        acento=AMBAR, tam_cuerpo=14)

texto(slide, 0.95, 4.55, 11.4, 1.2,
      "Había que resolver las dos cosas a la vez: que el dueño entre y que nadie más entre.",
      tam=21, color=VERDE, negrita=True, alineacion=PP_ALIGN.CENTER)

# ---------------------------------------------------------------------------
# 9. El pase firmado
# ---------------------------------------------------------------------------

slide = nueva_diapositiva()
titulo_diapositiva(slide, "La solución: un pase firmado", "Un enlace que no se puede fabricar", 9)

tarjeta(slide, 0.95, 1.9, 3.6, 2.6, "Firmado",
        ["Con la clave de la aplicación.",
         "Si alguien lo altera, deja de valer."], acento=VERDE, tam_cuerpo=13)

tarjeta(slide, 4.85, 1.9, 3.6, 2.6, "Con caducidad",
        ["8 horas de validez.",
         "Una dirección copiada deja de servir."], acento=VERDE, tam_cuerpo=13)

tarjeta(slide, 8.75, 1.9, 3.7, 2.6, "Atado al dueño",
        ["Va ligado al usuario y a su escritorio.",
         "El pase de uno no abre el de otro."], acento=VERDE, tam_cuerpo=13)

texto(slide, 0.95, 4.85, 11.4, 0.6, "Y viaja en la dirección, no en una cookie:",
      tam=18, color=BLANCO, negrita=True)

texto(slide, 0.95, 5.45, 11.4, 1.2,
      "…/contenedor/ventana?pase=firma&usuario=frank&contenedor=ct3",
      tam=17, color=VERDE)

# ---------------------------------------------------------------------------
# 10. El control de acceso
# ---------------------------------------------------------------------------

slide = nueva_diapositiva()
titulo_diapositiva(slide, "El control de acceso", "Cinco pasos antes de entregar un escritorio", 10)

pasos = [
    ("1", "El navegador pide el escritorio con su pase."),
    ("2", "Caddy añade la cabecera X-Container."),
    ("3", "Caddy pregunta a /internal/proxy-auth."),
    ("4", "La aplicación valida firma, fecha y dueño."),
    ("5", "Responde 200 o 403. Caddy entrega o niega."),
]

arriba = 1.95
for numero, descripcion in pasos:
    circulo = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(1.0), Inches(arriba), Inches(0.5), Inches(0.5))
    circulo.fill.solid()
    circulo.fill.fore_color.rgb = VERDE
    circulo.line.fill.background()
    marco = circulo.text_frame
    marco.text = numero
    marco.paragraphs[0].alignment = PP_ALIGN.CENTER
    run = marco.paragraphs[0].runs[0]
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = FONDO
    run.font.name = FUENTE

    texto(slide, 1.75, arriba - 0.02, 10.4, 0.6, descripcion, tam=17, color=BLANCO)
    arriba += 0.92

texto(slide, 1.0, 6.55, 11.4, 0.6,
      "El nombre del contenedor lo pone Caddy, no el navegador: no se puede falsear.",
      tam=13, color=GRIS)

# ---------------------------------------------------------------------------
# 11. Separador: moderacion
# ---------------------------------------------------------------------------

slide = nueva_diapositiva()
separador(slide, "Parte 2", "Moderación de imágenes",
          "NSFW-Sherlock revisa cada foto antes de publicarse.")

# ---------------------------------------------------------------------------
# 12. NSFW que es
# ---------------------------------------------------------------------------

slide = nueva_diapositiva()
titulo_diapositiva(slide, "NSFW-Sherlock", "Un servicio aparte, con una sola tarea", 12)

vinietas(slide, [
    ("Analiza una imagen y estima si es contenido no apto.", 0),
    ("Corre en su propio contenedor, sin publicar puertos.", 0),
    ("Solo se llega a él desde la red interna de Docker.", 0),
    ("La aplicación le habla por HTTP y le manda la imagen.", 0),
    ("La decisión de publicar o no es de la aplicación, no del servicio.", 0),
], arriba=1.9, tam=18, alto=3.4)

tarjeta(slide, 1.4, 4.75, 10.5, 1.5, "¿Por qué separarlo?",
        ["Se puede cambiar el modelo sin tocar la aplicación, y ajustar la sensibilidad sin reentrenar nada."],
        acento=VERDE, tam_cuerpo=14)

# ---------------------------------------------------------------------------
# 13. Flujo NSFW
# ---------------------------------------------------------------------------

slide = nueva_diapositiva()
titulo_diapositiva(slide, "El flujo de una revisión", "Qué ocurre cuando alguien sube una foto", 13)

cajas = [
    ("Sube la foto", "El usuario publica", VERDE),
    ("Temporal", "La aplicación la guarda", AZUL),
    ("Sherlock", "Devuelve puntuaciones", AMBAR),
    ("Umbral", "0 a 1, ajustable", AMBAR),
    ("Decisión", "Publica o rechaza", ROJO),
]

izq = 0.55
for titulo_caja, detalle, color in cajas:
    tarjeta(slide, izq, 2.2, 2.3, 1.9, titulo_caja, [detalle], acento=color, tam_titulo=14, tam_cuerpo=11)
    if izq < 9.0:
        flecha(slide, izq + 2.34, 2.95, 0.22, 0.3, color=GRIS)
    izq += 2.56

texto(slide, 0.95, 4.55, 11.4, 0.5, "Si supera el umbral:", tam=18, color=ROJO, negrita=True)
vinietas(slide, [
    ("La imagen se borra y no se publica.", 0),
    ("El usuario recibe un aviso explicando qué pasó.", 0),
    ("El administrador queda enterado para revisarlo.", 0),
], arriba=5.05, tam=14, alto=1.6)

# ---------------------------------------------------------------------------
# 14. Umbral
# ---------------------------------------------------------------------------

slide = nueva_diapositiva()
titulo_diapositiva(slide, "El umbral: dónde ponerlo", "La herramienta propone, la persona decide", 14)

filas = [
    ("0.90", "Muy permisivo", "Solo lo evidente", VERDE),
    ("0.75", "Equilibrado", "Recomendado en clase", AMBAR),
    ("0.60", "Estricto", "Etapas con reportes", ROJO),
]

arriba = 2.0
for valor, nombre, uso, color in filas:
    tarjeta(slide, 1.2, arriba, 3.0, 1.15, valor, [nombre], acento=color, tam_titulo=26, tam_cuerpo=12)
    texto(slide, 4.5, arriba + 0.35, 6.0, 0.6, uso, tam=16, color=GRIS)
    arriba += 1.45

texto(slide, 0.95, 6.45, 11.4, 0.6,
      "Es una revisión probabilística: puede equivocarse. Ante una duda, revisar a mano.",
      tam=14, color=AMBAR)

# ---------------------------------------------------------------------------
# 15. Separador: antivirus
# ---------------------------------------------------------------------------

slide = nueva_diapositiva()
separador(slide, "Parte 3", "Antivirus de archivos",
          "ClamAV revisa cada archivo antes de darlo por bueno.")

# ---------------------------------------------------------------------------
# 16. ClamAV
# ---------------------------------------------------------------------------

slide = nueva_diapositiva()
titulo_diapositiva(slide, "ClamAV: cómo se escanea", "Por socket, sin compartir archivos", 16)

vinietas(slide, [
    ("La subida llega por trozos; al terminar se juntan en un archivo temporal.", 0),
    ("Ese archivo se envía al antivirus con el comando INSTREAM.", 0),
    ("Se manda en bloques de 64 KB, sin copiar el archivo al otro contenedor.", 0),
    ("ClamAV responde: limpio, o con el nombre del virus.", 0),
    ("Solo si está limpio se mueve a su destino definitivo.", 0),
], arriba=1.9, tam=18, alto=3.4)

tarjeta(slide, 0.95, 5.05, 5.4, 1.3, "INSTREAM",
        ["Evita compartir el sistema de archivos entre contenedores."],
        acento=VERDE, tam_cuerpo=13)

tarjeta(slide, 6.9, 5.05, 5.5, 1.3, "Base de firmas",
        ["Unos 300 MB, guardados en un volumen para no repetir la descarga."],
        acento=AZUL, tam_cuerpo=13)

# ---------------------------------------------------------------------------
# 17. Fail open / closed
# ---------------------------------------------------------------------------

slide = nueva_diapositiva()
titulo_diapositiva(slide, "Si el antivirus no responde", "Una decisión de criterio, no técnica", 17)

tarjeta(slide, 1.1, 2.0, 5.3, 3.2, "Modo abierto",
        ["El archivo pasa.",
         "El incidente queda en el registro.",
         "Recomendado en un grupo de clase:",
         "un servicio caído no debe impedir subir apuntes."],
        acento=AMBAR, tam_cuerpo=14)

tarjeta(slide, 7.0, 2.0, 5.3, 3.2, "Modo cerrado",
        ["El archivo se rechaza.",
         "El usuario recibe el motivo.",
         "Recomendado cuando hay riesgo real",
         "de malware."],
        acento=ROJO, tam_cuerpo=14)

texto(slide, 0.95, 5.55, 11.4, 0.8,
      "Se configura con una variable de entorno, sin tocar código.",
      tam=17, color=GRIS, alineacion=PP_ALIGN.CENTER)

# ---------------------------------------------------------------------------
# 18. Subidas grandes
# ---------------------------------------------------------------------------

slide = nueva_diapositiva()
titulo_diapositiva(slide, "Subidas de hasta 5 GB", "Por trozos, con reanudación", 18)

vinietas(slide, [
    ("Los archivos se parten en trozos de 8 MB.", 0),
    ("Si se corta la conexión, el avance se conserva.", 0),
    ("Se puede reanudar donde se quedó.", 0),
    ("El botón Cancelar detiene la subida y no deja restos.", 0),
    ("Al terminar, el archivo completo pasa por el antivirus.", 0),
], arriba=1.9, tam=18, alto=3.6)

# ---------------------------------------------------------------------------
# 19. Diseno
# ---------------------------------------------------------------------------

slide = nueva_diapositiva()
titulo_diapositiva(slide, "El diseño: Redmond", "Cinco épocas de Windows, treinta esquemas", 19)

imagen_ajustada(slide, os.path.join(IMG, "04-temas-carrusel.png"), 0.9, 1.9, 5.2, 4.3)
imagen_ajustada(slide, os.path.join(IMG, "05-inicio-admin-calendario.png"), 6.4, 1.9, 6.2, 4.3)

texto(slide, 0.95, 6.4, 11.4, 0.6,
      "XP, Vista, 7, 8 y 10, con seis esquemas cada uno. El tema se recuerda por navegador.",
      tam=14, color=GRIS)

# ---------------------------------------------------------------------------
# 20. IA local
# ---------------------------------------------------------------------------

slide = nueva_diapositiva()
titulo_diapositiva(slide, "El asistente TutorUTT", "IA local, sin mandar nada a internet", 20)

vinietas(slide, [
    ("El modelo corre en el equipo del administrador, con Ollama.", 0),
    ("Ayuda con las materias de la carrera: explica, ejemplifica y propone ejercicios.", 0),
    ("Las preguntas no salen de la universidad.", 0),
    ("El panel muestra el uso de la tarjeta gráfica en vivo.", 0),
    ("Hoy está habilitado solo para el administrador.", 0),
], arriba=1.85, tam=17, alto=3.6)

imagen_ajustada(slide, os.path.join(IMG, "13-ia-tutorutt.png"), 9.2, 1.85, 3.4, 4.4)

# ---------------------------------------------------------------------------
# 21. Chat
# ---------------------------------------------------------------------------

slide = nueva_diapositiva()
titulo_diapositiva(slide, "El chat", "Conversación entre estudiantes", 21)

vinietas(slide, [
    ("Se abre desde cualquier página, y siempre en la lista de usuarios.", 0),
    ("Aviso de «está escribiendo» y marca de «visto».", 0),
    ("Contador de mensajes sin leer junto a cada persona.", 0),
    ("El «visto» solo se marca cuando de verdad se abre la conversación.", 0),
    ("Si estás leyendo, suena el aviso pero no salta la notificación.", 0),
    ("Los avisos se pausan cuando la pestaña está en segundo plano.", 0),
], arriba=1.85, tam=16, alto=4.3)

imagen_ajustada(slide, os.path.join(IMG, "12-contenedor-chat.png"), 7.4, 1.85, 5.2, 4.5)

# ---------------------------------------------------------------------------
# 22. Rendimiento
# ---------------------------------------------------------------------------

slide = nueva_diapositiva()
titulo_diapositiva(slide, "Hecho para equipos modestos", "Cada decisión pensada en un Core 2 Duo", 22)

tarjeta(slide, 0.9, 1.95, 3.75, 2.5, "Detección automática",
        ["Si el equipo tiene pocos núcleos o poca memoria, se activa el modo ligero."],
        acento=VERDE, tam_cuerpo=13)

tarjeta(slide, 4.8, 1.95, 3.75, 2.5, "Menos efectos",
        ["Se apagan desenfoques y fondos fijos: es lo que más cuesta dibujar."],
        acento=VERDE, tam_cuerpo=13)

tarjeta(slide, 8.7, 1.95, 3.75, 2.5, "Menos peticiones",
        ["Los avisos se pausan con la pestaña oculta.", "El sondeo es más espaciado."],
        acento=VERDE, tam_cuerpo=13)

vinietas(slide, [
    ("Fondos a resolución real de pantalla, sin estiramientos.", 0),
    ("La hoja del tema entra en el primer dibujado: sin parpadeo al recargar.", 0),
], arriba=4.85, tam=15, alto=1.6)

# ---------------------------------------------------------------------------
# 23. Seguridad
# ---------------------------------------------------------------------------

slide = nueva_diapositiva()
titulo_diapositiva(slide, "Seguridad en conjunto", "Capas, no una sola puerta", 23)

vinietas(slide, [
    ("HTTPS y cookies seguras en todo el sitio público.", 0),
    ("Contraseñas con hash; nunca en texto plano.", 0),
    ("Verificación en dos pasos con Google Authenticator.", 0),
    ("Permisos de administrador comprobados en el servidor, no en la interfaz.", 0),
    ("Límites de peticiones por usuario, no por dirección IP.", 0),
    ("Moderación de imágenes y antivirus en cada archivo.", 0),
    ("Pases firmados para los escritorios.", 0),
], arriba=1.9, tam=17, alto=4.6)

# ---------------------------------------------------------------------------
# 24. Demostracion
# ---------------------------------------------------------------------------

slide = nueva_diapositiva()
titulo_diapositiva(slide, "Así se ve", "Capturas del sistema en uso", 24)

imagen_ajustada(slide, os.path.join(IMG, "18-inicio-admin.png"), 0.7, 1.85, 6.0, 2.4)
imagen_ajustada(slide, os.path.join(IMG, "17-crear-publicacion.png"), 6.9, 1.85, 5.8, 2.4)
imagen_ajustada(slide, os.path.join(IMG, "12-contenedor-chat.png"), 0.7, 4.4, 6.0, 2.4)
imagen_ajustada(slide, os.path.join(IMG, "15-ventana-nueva.png"), 6.9, 4.4, 5.8, 2.4)

# ---------------------------------------------------------------------------
# 25. Cierre
# ---------------------------------------------------------------------------

slide = nueva_diapositiva(RGBColor(0x08, 0x14, 0x12))

banda = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.22), ALTO)
banda.fill.solid()
banda.fill.fore_color.rgb = VERDE
banda.line.fill.background()

texto(slide, 1.2, 1.55, 10.8, 0.5, "AWTY  ·  «are we there yet?»", tam=15, color=VERDE, negrita=True)
texto(slide, 1.2, 2.1, 10.8, 1.0, "¿Ya estamos ahí?", tam=54, color=BLANCO, negrita=True, fuente=FUENTE_TITULO)

texto(slide, 1.2, 3.35, 10.6, 2.2, [
    "Este proyecto se hizo para estudiantes que no tienen dinero suficiente",
    "para cambiar su equipo y necesitan herramientas demandantes.",
    "No es solo para hoy: lleva la filosofía de Linux,",
    "de la comunidad para la comunidad.",
], tam=18, color=GRIS, espaciado=6)

texto(slide, 1.2, 5.75, 10.8, 1.0, [
    "Francisco González  ·  Virthub 2.0",
    "Universidad Tecnológica de Tula-Tepeji",
], tam=14, color=VERDE, espaciado=2)

prs.save(SALIDA)
print(f"Guardado: {SALIDA}")

# Verificacion
comprobado = Presentation(SALIDA)
print(f"  diapositivas: {len(comprobado.slides)}")
print(f"  tamano: {comprobado.slide_width/914400:.2f} x {comprobado.slide_height/914400:.2f} pulgadas")

sin_texto = [i + 1 for i, s in enumerate(comprobado.slides)
             if not any(sh.has_text_frame and sh.text_frame.text.strip() for sh in s.shapes)]
print(f"  diapositivas sin texto: {sin_texto if sin_texto else 'ninguna'}")
