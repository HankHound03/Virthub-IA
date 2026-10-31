"""
Genera el Manual de Administrador de Virthub 2.0 (DOCX).
"""

import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from _kit_docx import (  # noqa: E402
    imagen,
    imagen_con_pie,
    indice,
    h1,
    h2,
    h3,
    nota,
    nuevo_documento,
    parrafo,
    pasos,
    portada,
    salto,
    tabla,
    vineta,
)

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SALIDA = os.path.join(RAIZ, "docs", "Manual-de-Administrador-Virthub-2.0.docx")


doc = nuevo_documento()

portada(
    doc,
    "Manual del Administrador",
    "Virthub 2.0",
    "Arquitectura, seguridad y mantenimiento\nUniversidad Tecnológica de Tula-Tepeji",
    imagen_portada=imagen("07-admin-usuarios.png"),
)

indice(
    doc,
    "Contenido",
    [
        (1, "1. Arquitectura general"),
        (1, "2. Puesta en marcha"),
        (1, "3. Variables de entorno"),
        (1, "4. El panel de administración"),
        (2, "4.1 Crear y gestionar usuarios"),
        (2, "4.2 Reportes del foro y sugerencias"),
        (1, "5. Moderación de imágenes: NSFW-Sherlock"),
        (2, "5.1 Cómo está montado"),
        (2, "5.2 El flujo de una revisión"),
        (2, "5.3 Ajustar el umbral"),
        (2, "5.4 Qué ve el usuario"),
        (1, "6. Antivirus: ClamAV"),
        (2, "6.1 Cómo se escanea"),
        (2, "6.2 Base de firmas y salud"),
        (2, "6.3 Fail-open y fail-closed"),
        (1, "7. Escritorios virtuales y tokenización"),
        (2, "7.1 Los contenedores Webtop"),
        (2, "7.2 El problema que había"),
        (2, "7.3 El pase firmado (DesktopPass)"),
        (2, "7.4 El control de acceso del proxy"),
        (2, "7.5 Cuándo aparece un 403"),
        (1, "8. Proxy y túnel"),
        (1, "9. La IA local"),
        (1, "10. Copias de seguridad"),
        (1, "11. Mantenimiento"),
        (1, "12. Seguridad"),
        (1, "13. Diagnóstico rápido"),
    ],
)

# --------------------------------------------------------------------------
# 1. Arquitectura
# --------------------------------------------------------------------------

h1(doc, "1. Arquitectura general")

parrafo(
    doc,
    "Virthub es una aplicación Laravel que se apoya en varios servicios independientes, todos "
    "dentro de Docker. Cada uno hace una cosa y se puede reiniciar sin tumbar los demás.",
)

tabla(
    doc,
    ["Servicio", "Imagen", "Función"],
    [
        ["app", "virthub-app", "La aplicación: Laravel 13, PHP 8.4, Nginx y PHP-FPM."],
        ["caddy", "caddy:2-alpine", "Proxy inverso. Reparte el tráfico y controla el acceso a los escritorios."],
        ["mariadb", "mariadb:11.4", "Base de datos: usuarios, publicaciones, reportes."],
        ["nsfw-sherlock", "m1chl/nsfw-sherlock", "Detector de contenido no apto para imágenes."],
        ["clamav", "clamav/clamav:stable", "Antivirus de los archivos subidos."],
        ["ct0 … ctN", "virthub-desktop:clon", "Escritorios Webtop: un contenedor por escritorio virtual."],
    ],
    anchos=[1.3, 1.9, 3.4],
)

parrafo(
    doc,
    "Los datos que no van a la base de datos viven en archivos JSON dentro de la aplicación: "
    "chat, amistades, avisos, perfiles y publicaciones del foro. Es una decisión deliberada: se "
    "pueden copiar con solo copiar la carpeta y no dependen del motor de base de datos.",
)

nota(
    doc,
    "Antes de tocar nada",
    "Los comandos de este manual se ejecutan en la raíz del proyecto. Producción y pruebas son "
    "dos proyectos de Docker distintos: no los mezcle.",
    tono="atencion",
)

# --------------------------------------------------------------------------
# 2. Puesta en marcha
# --------------------------------------------------------------------------

h1(doc, "2. Puesta en marcha")

parrafo(doc, "Producción:")

_codigo = (
    "docker compose -p virthub -f docker-compose.yml --env-file .env up -d\n"
    "docker compose -p virthub -f docker-compose.yml ps"
)
parrafo(doc, _codigo, negrita=False)

parrafo(doc, "Entorno de pruebas (puerto 8001):")

parrafo(
    doc,
    "docker compose -p virthub-test -f docker-compose.test.yml up -d",
)

tabla(
    doc,
    ["Situación", "Qué hacer"],
    [
        ["Cambió PHP, rutas o configuración", "Reconstruir la imagen: docker compose … build app y up -d."],
        ["Cambió solo CSS o JavaScript", "Se puede copiar al contenedor sin reconstruir."],
        ["Todo responde 500", "Ver logs de la aplicación y comprobar que la base de datos esté sana."],
    ],
    anchos=[2.6, 4.0],
)

nota(
    doc,
    "Producción no recarga PHP sola",
    "En producción, la caché de opcode no revisa si el archivo cambió. Después de tocar código PHP "
    "hay que reconstruir la imagen y recrear el contenedor; cambiar el archivo dentro del "
    "contenedor no surte efecto.",
)

# --------------------------------------------------------------------------
# 3. Variables
# --------------------------------------------------------------------------

h1(doc, "3. Variables de entorno")

tabla(
    doc,
    ["Variable", "Para qué sirve", "Valor típico"],
    [
        ["VIRTHUB_TEMAS", "Carrusel de temas Redmond.", "true"],
        ["VIRTHUB_NOTICIAS", "Gadget de noticias.", "false"],
        ["VIRTHUB_2FA", "Verificación en dos pasos.", "true en producción"],
        ["VIRTHUB_MODERACION", "Moderación NSFW.", "true"],
        ["VIRTHUB_MODERACION_UMBRAL", "Sensibilidad del filtro (0 a 1).", "0.75"],
        ["VIRTHUB_MODERACION_URL", "Dirección del servicio Sherlock.", "http://nsfw-sherlock:4000"],
        ["CLAMAV_ENABLED", "Antivirus de archivos.", "true"],
        ["OLLAMA_BASE_URL", "Servidor de la IA local.", "http://host.docker.internal:11434"],
        ["OLLAMA_MODEL", "Modelo de la IA.", "virthub-ia"],
        ["VIRTHUB_IA_MONITOR", "Monitor de GPU en el panel.", "true"],
        ["VIRTHUB_GPU_URL", "Puente de métricas de la GPU.", "http://host.docker.internal:11435"],
        ["VIRTHUB_VERSION", "Texto del pie de página.", "Virthub 2.0"],
    ],
    anchos=[1.9, 2.8, 1.9],
)

parrafo(
    doc,
    "El archivo .env no se versiona. Contiene la clave de la aplicación (APP_KEY), que es la que "
    "firma los pases de los escritorios: si se pierde, todos los pases dejan de valer.",
)

# --------------------------------------------------------------------------
# 4. Panel
# --------------------------------------------------------------------------

h1(doc, "4. El panel de administración")

parrafo(
    doc,
    "El botón «Admin» del encabezado abre el panel. Desde ahí se gestionan cuentas, se revisan "
    "reportes y se leen las sugerencias.",
)

imagen_con_pie(
    doc,
    imagen("07-admin-usuarios.png"),
    "Figura 1. Panel: creación de usuarios y listado de cuentas.",
)

h2(doc, "4.1 Crear y gestionar usuarios")

pasos(
    doc,
    [
        "Escriba el nombre visible y, opcionalmente, el nombre de usuario.",
        "Deje la contraseña vacía si quiere que se genere automáticamente.",
        "Elija el rol: user o admin.",
        "Pulse «Crear Usuario».",
    ],
)

parrafo(
    doc,
    "En la lista puede desactivar una cuenta (deja de poder entrar sin borrar su historial) o "
    "eliminarla. La desactivación es reversible; el borrado no.",
)

parrafo(
    doc,
    "También hay un apartado para cambiar la contraseña de cualquier usuario existente, útil "
    "cuando alguien la olvida y no tiene códigos de respaldo.",
)

h2(doc, "4.2 Reportes del foro y sugerencias")

parrafo(
    doc,
    "Los reportes llegan a la sección «Reportes del Foro». Cada uno indica la publicación, quién "
    "lo reportó y el motivo. Al revisarlo puede retirar la publicación; si lo hace, el autor "
    "recibe un aviso explicando el motivo.",
)

imagen_con_pie(
    doc,
    imagen("08-admin-reportes-sugerencias.png"),
    "Figura 2. Reportes del foro y sugerencias recibidas.",
)

parrafo(
    doc,
    "Las sugerencias se pueden enviar de forma identificada o anónima. Si son anónimas, no hay "
    "forma de saber quién las escribió: es intencional.",
)


# --------------------------------------------------------------------------
# 5. NSFW
# --------------------------------------------------------------------------

h1(doc, "5. Moderación de imágenes: NSFW-Sherlock")

parrafo(
    doc,
    "Sherlock es un servicio aparte que analiza una imagen y estima si es contenido no apto. La "
    "aplicación no incluye el modelo: le habla por HTTP y traduce su respuesta a una decisión.",
)

h2(doc, "5.1 Cómo está montado")

vineta(doc, "Imagen oficial m1chl/nsfw-sherlock, sin modificar.")
vineta(doc, "No publica puertos: solo se llega desde la red interna de Docker.")
vineta(doc, "La aplicación lo consulta en http://nsfw-sherlock:4000.")
vineta(doc, "Tiene comprobación de salud cada 15 segundos contra /ping.")

parrafo(
    doc,
    "El servicio recibe la imagen codificada en base64 y devuelve sus puntuaciones. La "
    "interpretación, el umbral y la decisión son de la aplicación, no del servicio: así se puede "
    "ajustar la sensibilidad sin reentrenar nada.",
)

h2(doc, "5.2 El flujo de una revisión")

tabla(
    doc,
    ["Paso", "Qué ocurre"],
    [
        ["1", "El usuario sube una foto al foro."],
        ["2", "La aplicación la guarda en temporal y llama a Sherlock."],
        ["3", "Sherlock responde con sus puntuaciones de contenido no apto."],
        ["4", "La aplicación compara con el umbral configurado."],
        ["5", "Si supera el umbral: se borra la imagen y se avisa al usuario y al administrador."],
        ["6", "Si no lo supera: la imagen se publica con normalidad."],
    ],
    anchos=[0.8, 5.8],
)

imagen_con_pie(
    doc,
    imagen("04-temas-carrusel.png"),
    "Figura 3. Las fotos del foro pasan por el filtro antes de publicarse.",
    ancho=3.4,
)

h2(doc, "5.3 Ajustar el umbral")

parrafo(
    doc,
    "El umbral vive en VIRTHUB_MODERACION_UMBRAL y va de 0 a 1. Más bajo significa más estricto: "
    "marca más imágenes, incluidas algunas normales. Más alto significa más permisivo.",
)

tabla(
    doc,
    ["Valor", "Comportamiento", "Cuándo usarlo"],
    [
        ["0.90", "Muy permisivo", "Casi solo contenido explícito evidente."],
        ["0.75", "Equilibrado (recomendado)", "Uso general en un grupo de clase."],
        ["0.60", "Estricto", "Etapas con muchos reportes."],
    ],
    anchos=[1.0, 2.6, 3.0],
)

parrafo(
    doc,
    "Es una revisión probabilística: puede equivocarse. Si alguien reporta un falso positivo, "
    "reviselo a mano antes de sancionar. La herramienta propone, la persona decide.",
)

h2(doc, "5.4 Qué ve el usuario")

parrafo(
    doc,
    "Si su foto se rechaza, el usuario recibe un aviso dentro de la aplicación: qué pasó, que fue "
    "una revisión automática y cómo apelar. No se le muestra la puntuación exacta, para no dar "
    "pistas de cómo evadir el filtro.",
)

nota(
    doc,
    "Si Sherlock se cae",
    "La aplicación puede optar por dejar pasar las imágenes (modo abierto) o bloquearlas (modo "
    "cerrado). Para un grupo de clase se recomienda el modo abierto más el reporte manual: un "
    "servicio caído no debe impedir publicar apuntes.",
)

# --------------------------------------------------------------------------
# 6. ClamAV
# --------------------------------------------------------------------------

h1(doc, "6. Antivirus: ClamAV")

parrafo(
    doc,
    "Todo archivo que se sube al foro pasa por ClamAV antes de darse por bueno. El antivirus corre "
    "en su propio contenedor y la aplicación le habla por socket.",
)

h2(doc, "6.1 Cómo se escanea")

vineta(doc, "La subida se hace por trozos; al terminar, los trozos se juntan en un archivo temporal.")
vineta(doc, "Ese archivo se envía a clamd con el comando INSTREAM, en bloques de 64 KB.")
vineta(doc, "clamd responde si está limpio o con el nombre del virus encontrado.")
vineta(doc, "Solo si está limpio se mueve a su destino definitivo.")
vineta(doc, "Si hay virus, el archivo se descarta y se avisa al usuario.")

parrafo(
    doc,
    "Se usa INSTREAM a propósito: evita compartir el sistema de archivos entre el contenedor de la "
    "aplicación y el del antivirus. El archivo sale de la aplicación, se analiza y nunca se "
    "ejecuta.",
)

h2(doc, "6.2 Base de firmas y salud")

parrafo(
    doc,
    "ClamAV descarga alrededor de 300 MB de firmas. Se guardan en un volumen para no repetir la "
    "descarga en cada arranque. La comprobación de salud usa clamdcheck.sh cada 30 segundos.",
)

h2(doc, "6.3 Fail-open y fail-closed")

tabla(
    doc,
    ["Modo", "Si el antivirus no responde", "Recomendado para"],
    [
        ["Abierto", "El archivo pasa y se anota el incidente en el registro.", "Un grupo de clase."],
        ["Cerrado", "El archivo se rechaza.", "Etapas con riesgo real de malware."],
    ],
    anchos=[1.4, 3.3, 1.9],
)

parrafo(
    doc,
    "Se puede desactivar del todo con CLAMAV_ENABLED=false, pero entonces no hay ninguna "
    "revisión de archivos. No lo haga en producción.",
)

# --------------------------------------------------------------------------
# 7. Contenedores y tokenizacion
# --------------------------------------------------------------------------

h1(doc, "7. Escritorios virtuales y tokenización")

parrafo(
    doc,
    "Cada escritorio es un contenedor Webtop independiente, con su propio nombre: ct0, ct1, ct2 y "
    "así sucesivamente. La aplicación decide cuál le toca a cada usuario.",
)

imagen_con_pie(
    doc,
    imagen("12-contenedor-chat.png"),
    "Figura 4. El administrador ve los botones para probar cada escritorio.",
    ancho=6.0,
)

h2(doc, "7.1 Los contenedores Webtop")

tabla(
    doc,
    ["Pieza", "Responsable"],
    [
        ["Contenedor del escritorio", "lscr.io/linuxserver/webtop con escritorio Ubuntu MATE."],
        ["Asignación a cada usuario", "La aplicación, según la cuenta."],
        ["Entrega al navegador", "Caddy, después de preguntar a la aplicación."],
        ["Comprobación de salud", "El propio contenedor publica su estado."],
    ],
    anchos=[2.2, 4.4],
)

h2(doc, "7.2 El problema que había")

parrafo(
    doc,
    "Antes, la comprobación de acceso dependía de la cookie de sesión. Pero la cookie no siempre "
    "viaja al subdominio del escritorio: el usuario, con sesión abierta, recibía un 403 al abrir "
    "su propio escritorio.",
)

parrafo(
    doc,
    "Y al revés: si se quitaba la comprobación, cualquiera con la dirección entraba. Había que "
    "resolver las dos cosas a la vez: que el dueño entre y que nadie más entre.",
)

h2(doc, "7.3 El pase firmado (DesktopPass)")

parrafo(
    doc,
    "La solución es un enlace firmado que viaja en la propia dirección, emitido por la aplicación "
    "al pulsar «Contenedor».",
)

tabla(
    doc,
    ["Propiedad", "Detalle"],
    [
        ["Firma", "Con la clave de la aplicación (APP_KEY): no se puede fabricar."],
        ["Caducidad", "8 horas (480 minutos). Una dirección copiada deja de servir."],
        ["Atadura", "Va atado al usuario y a su escritorio: el pase de uno no abre el de otro."],
        ["Dónde viaja", "En el parámetro «pase» de la dirección, no en una cookie."],
    ],
    anchos=[1.4, 5.2],
)

parrafo(
    doc,
    "Como no depende de cookies, funciona igual en el dominio principal y en el subdominio del "
    "escritorio. Y aporta algo que la cookie no daba: sin un pase emitido por la aplicación, la "
    "respuesta es 403, aunque alguien copie la dirección.",
)

h2(doc, "7.4 El control de acceso del proxy")

parrafo(
    doc,
    "Caddy, antes de entregar un escritorio, pregunta a la aplicación en /internal/proxy-auth. "
    "Esa dirección no es alcanzable desde fuera.",
)

tabla(
    doc,
    ["Paso", "Qué ocurre"],
    [
        ["1", "El navegador pide el escritorio con su pase en la dirección."],
        ["2", "Caddy añade la cabecera X-Container con el nombre del contenedor pedido."],
        ["3", "Caddy llama a /internal/proxy-auth con esa cabecera y el pase."],
        ["4", "La aplicación comprueba: ¿el pase es válido? ¿es suyo ese contenedor?"],
        ["5", "Responde 200 o 403. Caddy entrega el escritorio o lo niega."],
    ],
    anchos=[0.7, 5.9],
)

parrafo(
    doc,
    "El nombre del contenedor lo pone Caddy, no el cliente: no se puede falsear desde el "
    "navegador. La aplicación nunca devuelve un cuerpo útil en esa ruta, solo el código.",
)

h2(doc, "7.5 Cuándo aparece un 403")

tabla(
    doc,
    ["Motivo", "Cómo reconocerlo"],
    [
        ["El usuario no tiene escritorio asignado", "El botón Contenedor no aparece o falla siempre."],
        ["El pase caducó (más de 8 horas)", "Al volver a entrar desde Contenedor se arregla."],
        ["Se cambió APP_KEY", "Todos los pases dejan de valer hasta que se emitan de nuevo."],
        ["Se pide el escritorio de otro", "Solo pasa manipulando la dirección: es lo esperado."],
    ],
    anchos=[2.8, 3.8],
)


# --------------------------------------------------------------------------
# 8. Proxy
# --------------------------------------------------------------------------

h1(doc, "8. Proxy y túnel")

parrafo(
    doc,
    "Caddy escucha en el puerto 8080 en HTTP, porque el túnel ya termina el cifrado. Detrás del "
    "túnel, el sitio se sirve por HTTPS.",
)

vineta(doc, "El túnel de Cloudflare se ejecuta como servicio de Windows.")
vineta(doc, "Caddy reparte: la aplicación principal y los subdominios de cada escritorio.")
vineta(doc, "El puerto de la aplicación no se publica al exterior: solo Caddy lo alcanza.")
vineta(doc, "En pruebas, el puerto 8001 sí queda abierto en la red local para poder probar.")

# --------------------------------------------------------------------------
# 9. IA
# --------------------------------------------------------------------------

h1(doc, "9. La IA local")

parrafo(
    doc,
    "El asistente TutorUTT corre en Ollama, en el equipo del administrador, no en la nube. La "
    "aplicación lo alcanza por host.docker.internal:11434.",
)

tabla(
    doc,
    ["Pieza", "Detalle"],
    [
        ["Modelo", "virthub-ia, construido sobre granite4.2:3b (2.2 GB)."],
        ["Personalidad", "Archivo de instrucciones dentro del proyecto."],
        ["Acceso", "Solo el administrador; la ruta lo comprueba en el servidor."],
        ["Monitor", "Muestra uso y memoria de la GPU en el panel."],
        ["Puente de GPU", "Script de PowerShell que expone las métricas al contenedor."],
    ],
    anchos=[1.7, 4.9],
)

imagen_con_pie(
    doc,
    imagen("13-ia-tutorutt.png"),
    "Figura 5. TutorUTT con el monitor de recursos debajo.",
    ancho=3.0,
)

parrafo(
    doc,
    "Si el puente de la GPU no está levantado, el panel oculta esos datos y todo lo demás sigue "
    "funcionando. Sin Ollama, la pestaña IA responde con un aviso claro.",
)

# --------------------------------------------------------------------------
# 10. Copias
# --------------------------------------------------------------------------

h1(doc, "10. Copias de seguridad")

parrafo(
    doc,
    "Hay dos cosas que respaldar: la base de datos y los archivos de datos.",
)

tabla(
    doc,
    ["Qué", "Dónde", "Cómo"],
    [
        ["Base de datos", "Volumen de MariaDB", "Volcado con mariadb-dump."],
        ["Chat, amistades, avisos", "storage/app/data/*.json", "Copiar la carpeta."],
        ["Perfiles y publicaciones", "storage/app/data/*.json", "Copiar la carpeta."],
        ["Archivos subidos", "storage/app/uploads", "Copiar la carpeta."],
        ["Configuración", ".env (no versionado)", "Guardar aparte, con cuidado."],
    ],
    anchos=[1.8, 2.4, 2.4],
)

nota(
    doc,
    "La clave de la aplicación",
    "Si restaura una copia en otro equipo, use el mismo APP_KEY. Si cambia, los pases de los "
    "escritorios y las sesiones abiertas dejarán de valer.",
    tono="atencion",
)

# --------------------------------------------------------------------------
# 11. Mantenimiento
# --------------------------------------------------------------------------

h1(doc, "11. Mantenimiento")

h2(doc, "Archivos huérfanos del foro")

parrafo(
    doc,
    "Cuando se borra una publicación, sus archivos pueden quedarse en disco. El proyecto trae un "
    "comando para revisarlos y, si se pide, borrarlos:",
)

parrafo(doc, "php artisan foro:huerfanos          # solo informa")
parrafo(doc, "php artisan foro:huerfanos --borrar  # además borra")

parrafo(
    doc,
    "Ejecútelo primero sin --borrar y revise la lista. Después, con --borrar.",
)

h2(doc, "Registros")

vineta(doc, "Los errores de la aplicación quedan en el registro de Laravel.")
vineta(doc, "Los accesos a los escritorios se registran con el motivo de cada negativa.")
vineta(doc, "Los incidentes del antivirus y de la moderación también se anotan.")

h2(doc, "Actualizar el sitio")

pasos(
    doc,
    [
        "Haga una copia de los datos y de la base de datos.",
        "Traiga los cambios: git pull.",
        "Reconstruya la imagen: docker compose … build app.",
        "Recrete el contenedor: docker compose … up -d.",
        "Compruebe la salud y abra el sitio.",
    ],
)

# --------------------------------------------------------------------------
# 12. Seguridad
# --------------------------------------------------------------------------

h1(doc, "12. Seguridad")

tabla(
    doc,
    ["Tema", "Estado"],
    [
        ["HTTPS", "Todo el tráfico público va cifrado; las cookies son seguras."],
        ["Contraseñas", "Guardadas con hash, nunca en texto plano."],
        ["Verificación en dos pasos", "Disponible para todas las cuentas."],
        ["Permisos de administrador", "Comprobados en el servidor, no solo en la interfaz."],
        ["Límites de peticiones", "Por usuario en el chat, por dirección en el inicio de sesión."],
        ["Subidas", "Tope de tamaño, revisión de antivirus y moderación de imágenes."],
        ["Errores", "En producción no se muestran detalles internos."],
    ],
    anchos=[2.2, 4.4],
)

parrafo(
    doc,
    "Los límites de peticiones del chat cuentan por usuario y no por dirección IP, a propósito: en "
    "un aula, cien alumnos salen por la misma dirección y compartir el cupo dejaría a todos sin "
    "servicio.",
)

# --------------------------------------------------------------------------
# 13. Diagnostico
# --------------------------------------------------------------------------

h1(doc, "13. Diagnóstico rápido")

tabla(
    doc,
    ["Síntoma", "Qué revisar"],
    [
        ["El sitio no carga", "Túnel activo, Caddy arriba, aplicación sana."],
        ["Error 500 en una página", "Registro de la aplicación y espacio en disco."],
        ["El escritorio da 403", "Asignación del usuario y vigencia del pase."],
        ["El foro rechaza fotos buenas", "Subir el umbral de moderación."],
        ["Todas las subidas fallan", "Estado de ClamAV y de su base de firmas."],
        ["La IA no responde", "Ollama activo, modelo cargado y permisos de administrador."],
        ["El chat no avisa", "Permiso de notificaciones del navegador y sonido desbloqueado."],
        ["El monitor de GPU está vacío", "Puente de métricas levantado en el equipo."],
    ],
    anchos=[2.6, 4.0],
)

doc.save(SALIDA)
print(f"Guardado: {SALIDA}")

from docx import Document  # noqa: E402

comprobado = Document(SALIDA)
print(f"  parrafos: {len(comprobado.paragraphs)}")
print(f"  tablas: {len(comprobado.tables)}")
print(f"  imagenes: {len(comprobado.inline_shapes)}")
print(f"  titulos: {sum(1 for p in comprobado.paragraphs if p.style.name.startswith('Heading'))}")
